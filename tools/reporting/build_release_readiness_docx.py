#!/usr/bin/env python3
"""Build the ORCHESTRA-OS research and release-readiness DOCX brief."""

from __future__ import annotations

import argparse
import csv
import re
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "16324F"
TEAL = "137C8B"
PALE_TEAL = "DCEFF1"
AMBER = "D98E04"
PALE_AMBER = "FFF1CF"
RED = "A53A3A"
PALE_RED = "F8E3E3"
MID_GRAY = "596773"
LIGHT_GRAY = "EFF2F5"
WHITE = "FFFFFF"

METRIC_ORDER = ("S1_mean", "S2_mean", "S3_mean", "S4_mean", "Q_mean")


@dataclass(frozen=True)
class MetricSummary:
    mean: float
    lower: float
    upper: float
    n: int


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build a visual ORCHESTRA-OS release-readiness DOCX brief."
    )
    parser.add_argument(
        "--repository-root", type=Path, default=Path.cwd(), help="Repository root."
    )
    parser.add_argument(
        "--benchmark-dir",
        type=Path,
        default=Path("artifacts/test-results/2026-08-03/benchmark/authoritative-c"),
        help="Validated benchmark output containing processed/mode_summary.csv.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("output/doc/ORCHESTRA-OS_Release_Readiness_Plan.docx"),
        help="DOCX output path, relative to the repository root unless absolute.",
    )
    return parser.parse_args()


def sanitize_text(value: str) -> str:
    """Use portable punctuation while retaining the source meaning."""
    replacements = {
        "\u2010": "-",
        "\u2011": "-",
        "\u2012": "-",
        "\u2013": "-",
        "\u2014": " - ",
        "\u2212": "-",
        "\u00a0": " ",
    }
    for source, target in replacements.items():
        value = value.replace(source, target)
    return value


def shade_cell(cell: object, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()  # type: ignore[attr-defined]
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell: object, top: int = 80, start: int = 100,
                     bottom: int = 80, end: int = 100) -> None:
    tc = cell._tc  # type: ignore[attr-defined]
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row: object) -> None:
    tr_pr = row._tr.get_or_add_trPr()  # type: ignore[attr-defined]
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_page_number(paragraph: object) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend((field_begin, instruction, field_end))


def configure_document(document: DocumentType) -> None:
    document.core_properties.title = "ORCHESTRA-OS Research Evidence and Release-Readiness Plan"
    document.core_properties.subject = "Userspace benchmark visualization and WP1-WP10 release gates"
    document.core_properties.author = "ORCHESTRA-OS research program"
    document.core_properties.keywords = "ORCHESTRA-OS, scheduler, research, benchmark, release readiness"
    fixed_time = datetime(2026, 8, 3, 0, 0, 0, tzinfo=timezone.utc)
    document.core_properties.created = fixed_time
    document.core_properties.modified = fixed_time

    section = document.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.63)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.3)
    normal.font.color.rgb = RGBColor.from_string("263746")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color, before, after in (
        ("Title", 28, NAVY, 0, 8),
        ("Subtitle", 13, MID_GRAY, 0, 10),
        ("Heading 1", 19, NAVY, 12, 7),
        ("Heading 2", 13.5, TEAL, 10, 5),
        ("Heading 3", 10.5, NAVY, 8, 3),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Subtitle" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Figure Caption" not in styles:
        caption = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Figure Caption"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(8)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MID_GRAY)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)

    header = section.header.paragraphs[0]
    header.text = "ORCHESTRA-OS  |  RESEARCH EVIDENCE AND RELEASE READINESS"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    add_page_number(section.footer.paragraphs[0])


def add_rich_text(paragraph: object, text: str) -> None:
    """Render the small inline Markdown subset used by the plan."""
    text = sanitize_text(text)
    token_pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^]]+\]\([^)]+\))")
    cursor = 0
    for match in token_pattern.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Liberation Mono"
            run.font.size = Pt(8.3)
            run.font.color.rgb = RGBColor.from_string(NAVY)
        else:
            link_match = re.fullmatch(r"\[([^]]+)\]\(([^)]+)\)", token)
            label = link_match.group(1) if link_match else token
            target = link_match.group(2) if link_match else ""
            run = paragraph.add_run(f"{label} ({target})" if target else label)
            run.font.color.rgb = RGBColor.from_string(TEAL)
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def add_callout(document: DocumentType, title: str, body: str,
                *, fill: str, accent: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, top=130, start=170, bottom=130, end=170)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(3)
    title_run = paragraph.add_run(sanitize_text(title) + "\n")
    title_run.bold = True
    title_run.font.size = Pt(10.5)
    title_run.font.color.rgb = RGBColor.from_string(accent)
    body_run = paragraph.add_run(sanitize_text(body))
    body_run.font.size = Pt(9)
    body_run.font.color.rgb = RGBColor.from_string("263746")
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def style_table(table: object, *, first_column_bold: bool = False) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    if table.rows:
        set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                shade_cell(cell, NAVY)
            elif row_index % 2 == 0:
                shade_cell(cell, LIGHT_GRAY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(1.5)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(7.8)
                    if row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor.from_string(WHITE)
                    elif first_column_bold and column_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor.from_string(NAVY)


def add_table_from_rows(
    document: DocumentType,
    headers: Iterable[str],
    rows: Iterable[Iterable[str]],
    *,
    first_column_bold: bool = False,
) -> None:
    header_list = [sanitize_text(item) for item in headers]
    table = document.add_table(rows=1, cols=len(header_list))
    for index, value in enumerate(header_list):
        table.cell(0, index).text = value
    for source_row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(source_row):
            cells[index].text = sanitize_text(str(value))
    style_table(table, first_column_bold=first_column_bold)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def load_metrics(path: Path) -> dict[str, dict[str, MetricSummary]]:
    required = {"mode", "metric", "n", "mean", "ci95_lower", "ci95_upper"}
    result: dict[str, dict[str, MetricSummary]] = {}
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None or not required.issubset(reader.fieldnames):
            raise ValueError(f"missing required columns in {path}")
        for row in reader:
            metric = row["metric"]
            if metric not in METRIC_ORDER:
                continue
            mode = row["mode"]
            result.setdefault(mode, {})[metric] = MetricSummary(
                mean=float(row["mean"]),
                lower=float(row["ci95_lower"]),
                upper=float(row["ci95_upper"]),
                n=int(row["n"]),
            )
    for mode in ("baseline", "orchestra"):
        if set(result.get(mode, {})) != set(METRIC_ORDER):
            raise ValueError(f"incomplete coordination metrics for {mode}")
    return result


def add_cover(document: DocumentType) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(50)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run("ORCHESTRA-OS")
    run.font.name = "Aptos Display"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(TEAL)

    title = document.add_paragraph(style="Title")
    title.add_run("Research evidence and\nrelease-readiness plan")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run("Visual benchmark brief | Evidence-gated WP1-WP10 roadmap | 3 August 2026")

    add_callout(
        document,
        "Current claim class: USERSPACE-VALIDATED",
        "Selected userspace WP2-WP6 mechanics and a bounded six-run smoke campaign are available. "
        "Kernel-prototyped, Experimentally validated, and Deployment-ready claims are not supported.",
        fill=PALE_TEAL,
        accent=TEAL,
    )

    document.add_paragraph()
    document.add_heading("Decision in one sentence", level=2)
    paragraph = document.add_paragraph()
    paragraph.add_run(
        "Do not relabel the prototype as a release. Build and pass the kernel, trusted-loop, "
        "controlled-evidence, scale, resilience, and operational gates in dependency order; only "
        "then hold a formal deployment-readiness review."
    )

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_executive_snapshot(
    document: DocumentType, metrics: dict[str, dict[str, MetricSummary]]
) -> None:
    document.add_heading("1. Executive evidence snapshot", level=1)
    add_table_from_rows(
        document,
        ("Evidence item", "Observed result", "Interpretation boundary"),
        (
            (
                "Campaign validity",
                "6/6 invocations valid; 3 independent runs per mode; 30 post-warm-up rows per run",
                "Bounded smoke evidence, not general performance evidence",
            ),
            (
                "ORCHESTRA coordination",
                "Q 0.655; S1 0.629; S2 0.767; S3 0.716; S4 0.600",
                "Descriptive run means with n=3; intervals remain wide",
            ),
            (
                "Prediction path",
                "Prediction used for 95.6% of post-warm-up decisions on average",
                "Calibration was not independently held out",
            ),
            (
                "Integrity and cadence",
                "0 rejected frames, 0 publisher deadline misses; 2 controller updates per ORCHESTRA run",
                "Normal-path smoke only; not production security or reliability proof",
            ),
            (
                "Dispatch authority",
                "Linux CFS/EEVDF remained in control",
                "Userspace actions are approximations, not kernel action semantics",
            ),
        ),
        first_column_bold=True,
    )

    document.add_heading("Run-level coordination summary", level=2)
    rows: list[tuple[str, ...]] = []
    labels = {"baseline": "Observed-state reference", "orchestra": "ORCHESTRA"}
    for mode in ("baseline", "orchestra"):
        values = metrics[mode]
        row = [labels[mode]]
        for metric in METRIC_ORDER:
            summary = values[metric]
            row.append(f"{summary.mean:.3f} [{summary.lower:.3f}, {summary.upper:.3f}]")
        rows.append(tuple(row))
    add_table_from_rows(
        document,
        ("Mode", "S1 [95% CI]", "S2 [95% CI]", "S3 [95% CI]", "S4 [95% CI]", "Q [95% CI]"),
        rows,
        first_column_bold=True,
    )
    add_callout(
        document,
        "Baseline interpretation",
        "The observed-state reference scores 1.0 by construction under this contract. The modes are "
        "unpaired and alter the endogenous host signal, so subtraction, speedup, treatment effect, and "
        "scheduler-superiority claims are not valid.",
        fill=PALE_AMBER,
        accent=AMBER,
    )


def add_picture_if_present(
    document: DocumentType, path: Path, caption: str, *, width: float = 6.9
) -> bool:
    if not path.is_file():
        return False
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_paragraph = document.add_paragraph(style="Figure Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.add_run(sanitize_text(caption))
    return True


def select_figure(figure_dir: Path, terms: tuple[str, ...]) -> Path | None:
    candidates = sorted(figure_dir.glob("*.png"))
    for candidate in candidates:
        normalized = candidate.stem.lower().replace("-", "_")
        if all(term in normalized for term in terms):
            return candidate
    return None


def add_benchmark_figures(document: DocumentType, repository_root: Path) -> int:
    document.add_heading("2. Visual benchmark results", level=1)
    figure_dir = repository_root / "docs/experiments/figures/paper_cpu_smoke_v1_2026-08-03"
    figure_specs = (
        (
            ("coordination",),
            "Figure 1. Coordination metric means and 95% Student-t intervals across three independent runs per mode. The reference is a contract check, not a causal comparator.",
        ),
        (
            ("coordination", "per", "run"),
            "Figure 2. Per-run S1-S4 and Q values. Points expose run-to-run variability that a single aggregate would hide.",
        ),
        (
            ("action",),
            "Figure 3. Canonical action mix by run. RUN, SLEEP, MIGRATE, THROTTLE, and YIELD remain userspace approximations here.",
        ),
    )
    count = 0
    for terms, caption in figure_specs:
        path = select_figure(figure_dir, terms)
        if path is not None and add_picture_if_present(document, path, caption):
            count += 1
    if count == 0:
        add_callout(
            document,
            "Figure generation required",
            "Run tools/plotting/plot_paper_cpu_smoke.py against the authoritative benchmark directory before rebuilding this document.",
            fill=PALE_RED,
            accent=RED,
        )
    return count


def split_table_row(line: str) -> list[str]:
    return [sanitize_text(cell.strip()) for cell in line.strip().strip("|").split("|")]


def is_alignment_row(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def add_markdown_plan(document: DocumentType, path: Path) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    index = 0
    first_h1_skipped = False
    in_fence = False
    fence_kind = ""
    code_lines: list[str] = []
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        if not paragraph_lines:
            return
        paragraph = document.add_paragraph()
        add_rich_text(paragraph, " ".join(part.strip() for part in paragraph_lines))
        paragraph_lines.clear()

    while index < len(lines):
        raw_line = lines[index]
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            if not in_fence:
                in_fence = True
                fence_kind = stripped[3:].strip().lower()
                code_lines = []
            else:
                if fence_kind != "mermaid" and code_lines:
                    paragraph = document.add_paragraph()
                    run = paragraph.add_run("\n".join(sanitize_text(item) for item in code_lines))
                    run.font.name = "Liberation Mono"
                    run.font.size = Pt(7.7)
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), LIGHT_GRAY)
                    paragraph._p.get_or_add_pPr().append(shade)
                in_fence = False
                fence_kind = ""
                code_lines = []
            index += 1
            continue
        if in_fence:
            code_lines.append(line)
            index += 1
            continue

        if not stripped:
            flush_paragraph()
            index += 1
            continue

        heading_match = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading_match:
            flush_paragraph()
            level = len(heading_match.group(1))
            heading_text = sanitize_text(heading_match.group(2))
            if level == 1 and not first_h1_skipped:
                first_h1_skipped = True
                index += 1
                continue
            if level >= 2:
                heading_text = re.sub(r"^\d+\.\s*", "", heading_text)
            mapped_level = min(max(level, 1), 3)
            document.add_heading(heading_text, level=mapped_level)
            index += 1
            continue

        if "|" in stripped and index + 1 < len(lines) and is_alignment_row(lines[index + 1]):
            flush_paragraph()
            headers = split_table_row(stripped)
            index += 2
            rows: list[list[str]] = []
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                row = split_table_row(lines[index])
                if len(row) == len(headers):
                    rows.append(row)
                index += 1
            add_table_from_rows(document, headers, rows, first_column_bold=True)
            continue

        bullet_match = re.match(r"^\s*[-*]\s+(.+)$", line)
        numbered_match = re.match(r"^\s*\d+[.)]\s+(.+)$", line)
        if bullet_match or numbered_match:
            flush_paragraph()
            paragraph = document.add_paragraph(
                style="List Bullet" if bullet_match else "List Number"
            )
            add_rich_text(paragraph, (bullet_match or numbered_match).group(1))
            index += 1
            continue

        if stripped.startswith("!["):
            index += 1
            continue

        paragraph_lines.append(stripped)
        index += 1
    flush_paragraph()


def add_release_plan(document: DocumentType, repository_root: Path) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_heading("3. Evidence-gated path to release", level=1)
    add_picture_if_present(
        document,
        repository_root / "docs/operations/figures/release_readiness_path.png",
        "Figure 4. Dependency path from current userspace evidence through WP1-WP10. Boxes are gates, not a schedule or percent-complete estimate.",
    )
    plan_path = repository_root / "docs/operations/release-readiness-plan.md"
    if plan_path.is_file():
        add_markdown_plan(document, plan_path)
    else:
        add_callout(
            document,
            "Roadmap source missing",
            f"Expected {plan_path.relative_to(repository_root)}.",
            fill=PALE_RED,
            accent=RED,
        )


def add_evidence_appendix(document: DocumentType) -> None:
    document.add_heading("4. Evidence identity and document limits", level=1)
    add_table_from_rows(
        document,
        ("Artifact", "SHA-256"),
        (
            ("C source", "d5ec68c28a2cfcc20475dba44cb4b3c36b46b08acb4206647b874d2e2edf8713"),
            ("Benchmark binary", "f2f14020ef6f4caf7cc1f9dda3b392451367209b9d9f79e27415680c5e678b4b"),
            ("Benchmark harness", "3c07263066118a7f08c9e10ae80c488bb8fae680db08b31cba6b7154aec193d6"),
            ("Manifest", "40844535648f4a7dc25131141035d9cfb88886c20f0f8b9fb4bbbc8ffef02b5a"),
            ("Metrics schema", "04cb97e55d66bcf3e03e810372865c25368c89abb5de4e09e43d952d6f2d09c2"),
        ),
        first_column_bold=True,
    )
    add_callout(
        document,
        "Document limit",
        "This brief is a planning and evidence-communication artifact. It does not approve a release, "
        "replace an ADR, satisfy a work-package exit gate, or strengthen the underlying research claim.",
        fill=LIGHT_GRAY,
        accent=NAVY,
    )


def canonicalize_docx_archive(path: Path) -> None:
    """Normalize ZIP member timestamps so identical inputs are byte-stable."""
    temporary = path.with_suffix(path.suffix + ".canonicalizing")
    fixed_zip_time = (2026, 8, 3, 0, 0, 0)
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temporary, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as destination:
        for source_info in source.infolist():
            payload = source.read(source_info.filename)
            target_info = zipfile.ZipInfo(source_info.filename, date_time=fixed_zip_time)
            target_info.compress_type = zipfile.ZIP_DEFLATED
            target_info.comment = source_info.comment
            target_info.extra = source_info.extra
            target_info.external_attr = source_info.external_attr
            target_info.internal_attr = source_info.internal_attr
            target_info.create_system = source_info.create_system
            destination.writestr(target_info, payload)
    temporary.replace(path)


def build(repository_root: Path, benchmark_dir: Path, output: Path) -> tuple[int, int]:
    repository_root = repository_root.resolve()
    if not benchmark_dir.is_absolute():
        benchmark_dir = repository_root / benchmark_dir
    benchmark_dir = benchmark_dir.resolve()
    output = output if output.is_absolute() else repository_root / output
    mode_summary = benchmark_dir / "processed/mode_summary.csv"
    if not mode_summary.is_file():
        raise FileNotFoundError(f"authoritative mode summary not found: {mode_summary}")

    metrics = load_metrics(mode_summary)
    document = Document()
    configure_document(document)
    add_cover(document)
    add_executive_snapshot(document, metrics)
    figure_count = add_benchmark_figures(document, repository_root)
    add_release_plan(document, repository_root)
    add_evidence_appendix(document)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    canonicalize_docx_archive(output)

    verification = Document(output)
    if figure_count < 3:
        raise RuntimeError(f"expected at least 3 benchmark figures, found {figure_count}")
    if len(verification.inline_shapes) < 4:
        raise RuntimeError("generated document is missing expected visual figures")
    if len(verification.tables) < 5:
        raise RuntimeError("generated document is missing expected evidence tables")
    return len(verification.paragraphs), len(verification.tables)


def main() -> int:
    args = parse_args()
    paragraphs, tables = build(args.repository_root, args.benchmark_dir, args.output)
    output = args.output if args.output.is_absolute() else args.repository_root / args.output
    print(f"wrote {output.resolve()}")
    print(f"validated paragraphs={paragraphs} tables={tables}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
